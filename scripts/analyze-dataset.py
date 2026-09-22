"""Audit the strawberry RGB-D dataset and render publication-style web assets.

The script is intentionally read-only with respect to the dataset. It writes a
machine-readable summary and three deterministic visualizations to src/assets.
"""

from __future__ import annotations

import argparse
import json
import math
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from statistics import median

import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


KEYPOINTS = ("Cut", "Top", "Bottom", "Left", "Right")
MODALITIES = ("rgb", "label", "depth_aligned", "depth_aligned_pp", "depth_raw")
COLORS = {
    "navy": "#17364d",
    "blue": "#2b6172",
    "red": "#b33d4c",
    "green": "#547a63",
    "gold": "#c58b30",
    "violet": "#76578f",
    "cyan": "#2f8191",
    "ink": "#1b2731",
    "muted": "#5f6c76",
    "line": "#d5dde3",
    "paper": "#ffffff",
    "panel": "#f3f6f8",
}
KEYPOINT_COLORS = {
    "Cut": COLORS["gold"],
    "Top": COLORS["cyan"],
    "Bottom": COLORS["green"],
    "Left": COLORS["violet"],
    "Right": COLORS["red"],
}
ANNOTATION_KEYPOINT_RADIUS = 4
ANNOTATION_BOX_WIDTH = 2


def collection_sort_key(path: Path) -> int:
    digits = "".join(ch for ch in path.name if ch.isdigit())
    return int(digits or 0)


def font(size: int, bold: bool = False, mono: bool = False) -> ImageFont.FreeTypeFont:
    if mono:
        candidates = [
            Path("C:/Windows/Fonts/consolab.ttf" if bold else "C:/Windows/Fonts/consola.ttf"),
            Path("C:/Windows/Fonts/courbd.ttf" if bold else "C:/Windows/Fonts/cour.ttf"),
        ]
    else:
        candidates = [
            Path("C:/Windows/Fonts/seguisb.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf"),
            Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default(size=size)


def stems_in(directory: Path, suffix: str) -> set[str]:
    if not directory.exists():
        return set()
    return {path.stem for path in directory.glob(f"*{suffix}")}


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def value_sort_key(item: tuple[str, int]) -> tuple[int, str]:
    value = item[0]
    try:
        return (0, f"{int(value):08d}")
    except ValueError:
        return (1, value)


def percent(part: int, whole: int) -> float:
    return round(100.0 * part / whole, 2) if whole else 0.0


def discover_partitions(dataset_root: Path) -> tuple[str, list[dict]]:
    collections = sorted(
        [path for path in dataset_root.glob("data_collection*") if path.is_dir()],
        key=collection_sort_key,
    )
    if collections:
        return "collections", [
            {
                "name": collection.name,
                "label": collection.name.replace("data_collection", "C"),
                "label_dir": collection / "label",
                "rgb_dir": collection / "rgb",
                "depth_aligned_dir": collection / "depth_aligned",
                "depth_aligned_pp_dir": collection / "depth_aligned_pp",
                "depth_raw_dir": collection / "depth_raw",
                "calibration_root": collection,
            }
            for collection in collections
        ]

    split_names = [name for name in ("train", "val", "test") if (dataset_root / name).is_dir()]
    if split_names:
        return "train_val", [
            {
                "name": split_name,
                "label": split_name.capitalize(),
                "label_dir": dataset_root / split_name / "label",
                "rgb_dir": dataset_root / split_name / "image" / "rgb",
                "depth_aligned_dir": dataset_root / split_name / "image" / "depth_aligned",
                "depth_aligned_pp_dir": dataset_root / split_name / "image" / "depth_aligned_pp",
                "depth_raw_dir": dataset_root / split_name / "image" / "depth_raw",
                "calibration_root": dataset_root / "camera_calibration",
            }
            for split_name in split_names
        ]
    raise FileNotFoundError(f"No supported collection or train/val layout found under {dataset_root}")


def load_calibration_mapping(dataset_root: Path) -> dict[tuple[str, str], Path]:
    workbook_path = dataset_root / "image_camera_calibration_mapping.xlsx"
    if not workbook_path.exists():
        return {}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    indexes = {header: index for index, header in enumerate(headers)}
    required = {"split", "image_name", "calibration_file"}
    if not required.issubset(indexes):
        raise ValueError(f"Calibration mapping is missing required columns: {sorted(required - set(indexes))}")
    mapping: dict[tuple[str, str], Path] = {}
    for row in rows:
        split = str(row[indexes["split"]] or "").strip()
        image_name = str(row[indexes["image_name"]] or "").strip()
        calibration_file = str(row[indexes["calibration_file"]] or "").strip()
        if split and image_name and calibration_file:
            mapping[(split, Path(image_name).stem)] = dataset_root / Path(calibration_file)
    workbook.close()
    return mapping


def scan_dataset(dataset_root: Path, preferred_sample: str | None = None) -> tuple[dict, dict]:
    layout, partitions = discover_partitions(dataset_root)
    calibration_mapping = load_calibration_mapping(dataset_root)
    preferred_partition: str | None = None
    preferred_stem: str | None = None
    if preferred_sample:
        sample_parts = preferred_sample.replace("\\", "/").strip("/").split("/", 1)
        if len(sample_parts) != 2 or not all(sample_parts):
            raise ValueError("Representative sample must use the form 'partition/stem', for example 'train/0238'.")
        preferred_partition, preferred_stem = sample_parts
    overall = Counter()
    ripeness = Counter()
    occlusion = Counter()
    visibility = Counter()
    image_sizes = Counter()
    partition_rows: list[dict] = []
    candidate: dict | None = None
    camera_serials: set[str] = set()
    device_names: set[str] = set()

    calibration_roots = {partition["calibration_root"] for partition in partitions}
    for calibration_root in calibration_roots:
        for calibration_path in calibration_root.glob("camera_calibration*.json"):
            if calibration_path.name == "camera_calibration_mapping.json":
                continue
            try:
                calibration = read_json(calibration_path)
            except (json.JSONDecodeError, OSError):
                continue
            device = calibration.get("device", {})
            serial = str(device.get("serial_number", "")).strip()
            name = str(device.get("name", "")).strip()
            if serial:
                camera_serials.add(serial)
            if name:
                device_names.add(name)

    for partition in partitions:
        modality_stems = {
            "rgb": stems_in(partition["rgb_dir"], ".png"),
            "label": stems_in(partition["label_dir"], ".json"),
            "depth_aligned": stems_in(partition["depth_aligned_dir"], ".png"),
            "depth_aligned_pp": stems_in(partition["depth_aligned_pp_dir"], ".png"),
            "depth_raw": stems_in(partition["depth_raw_dir"], ".png"),
        }
        paired = set.intersection(*(modality_stems[name] for name in MODALITIES))
        partition_counts = Counter()
        partition_ripeness = Counter()
        partition_occlusion = Counter()
        partition_visibility = Counter()

        for label_path in sorted(partition["label_dir"].glob("*.json")):
            record = read_json(label_path)
            width = int(record.get("imageWidth") or 0)
            height = int(record.get("imageHeight") or 0)
            image_sizes[f"{width}x{height}"] += 1
            shapes = record.get("shapes", [])
            boxes: dict[object, dict] = {}
            points_by_group: dict[object, list[dict]] = defaultdict(list)

            for shape in shapes:
                shape_type = shape.get("shape_type")
                label = shape.get("label")
                group_id = shape.get("group_id")
                if shape_type == "rectangle" or label == "strawberry":
                    partition_counts["instances"] += 1
                    overall["instances"] += 1
                    boxes[group_id] = shape
                    ripeness_value = str(shape.get("ripeness", "missing"))
                    occlusion_value = str(shape.get("occlusion", "missing"))
                    ripeness[ripeness_value] += 1
                    occlusion[occlusion_value] += 1
                    partition_ripeness[ripeness_value] += 1
                    partition_occlusion[occlusion_value] += 1
                    coordinates = shape.get("points", [])
                    valid_box = False
                    if len(coordinates) == 2:
                        (x1, y1), (x2, y2) = coordinates
                        values = (x1, y1, x2, y2)
                        if all(isinstance(value, (int, float)) and math.isfinite(value) for value in values):
                            valid_box = abs(x2 - x1) > 0 and abs(y2 - y1) > 0
                    if not valid_box:
                        partition_counts["invalid_boxes"] += 1
                        overall["invalid_boxes"] += 1
                elif shape_type == "point" or label in KEYPOINTS:
                    partition_counts["keypoints"] += 1
                    overall["keypoints"] += 1
                    visibility_value = str(shape.get("visibility", "missing"))
                    visibility[visibility_value] += 1
                    partition_visibility[visibility_value] += 1
                    points_by_group[group_id].append(shape)

            complete_groups = 0
            for group_id in boxes:
                labels = Counter(point.get("label") for point in points_by_group.get(group_id, []))
                if all(labels[name] == 1 for name in KEYPOINTS) and sum(labels.values()) == len(KEYPOINTS):
                    complete_groups += 1
            unmatched_points = sum(
                len(group_points) for group_id, group_points in points_by_group.items() if group_id not in boxes
            )
            partition_counts["complete_groups"] += complete_groups
            partition_counts["unmatched_points"] += unmatched_points
            overall["complete_groups"] += complete_groups
            overall["unmatched_points"] += unmatched_points

            n_boxes = len(boxes)
            record_ripeness = {str(shape.get("ripeness", "missing")) for shape in boxes.values()}
            record_occlusion = {str(shape.get("occlusion", "missing")) for shape in boxes.values()}
            complete = n_boxes > 0 and complete_groups == n_boxes
            has_all_modalities = label_path.stem in paired
            calibration_path = calibration_mapping.get((partition["name"], label_path.stem))
            if calibration_path and calibration_path.exists():
                overall["calibration_mapped_samples"] += 1
            score = (
                (40 if 5 <= n_boxes <= 12 else max(0, 20 - abs(n_boxes - 8)))
                + (18 if len(record_ripeness) > 1 else 0)
                + (18 if len(record_occlusion) > 1 else 0)
                + (18 if complete else 0)
                + (12 if has_all_modalities else 0)
                + (10 if calibration_path and calibration_path.exists() else 0)
                + min(n_boxes, 12)
                + (
                    10_000
                    if partition["name"] == preferred_partition and label_path.stem == preferred_stem
                    else 0
                )
            )
            if candidate is None or score > candidate["score"]:
                candidate = {
                    "score": score,
                    "partition": partition["name"],
                    "partition_label": partition["label"],
                    "stem": label_path.stem,
                    "record": record,
                    "instances": n_boxes,
                    "complete": complete,
                    "rgb_path": partition["rgb_dir"] / f"{label_path.stem}.png",
                    "depth_pp_path": partition["depth_aligned_pp_dir"] / f"{label_path.stem}.png",
                    "calibration_path": calibration_path,
                    "legacy_calibration_root": partition["calibration_root"],
                }

        row = {
            "partition": partition["name"],
            "label": partition["label"],
            "samples": len(modality_stems["label"]),
            "instances": partition_counts["instances"],
            "keypoints": partition_counts["keypoints"],
            "schema_complete_instances": partition_counts["complete_groups"],
            "invalid_boxes": partition_counts["invalid_boxes"],
            "unmatched_keypoints": partition_counts["unmatched_points"],
            "ripeness": dict(sorted(partition_ripeness.items(), key=value_sort_key)),
            "occlusion": dict(sorted(partition_occlusion.items(), key=value_sort_key)),
            "keypoint_visibility": dict(sorted(partition_visibility.items(), key=value_sort_key)),
            "modalities": {name: len(modality_stems[name]) for name in MODALITIES},
            "fully_paired_samples": len(paired),
        }
        partition_rows.append(row)
        overall["fully_paired_samples"] += len(paired)
        for name in MODALITIES:
            overall[f"{name}_files"] += len(modality_stems[name])

    extension_counts = Counter()
    total_bytes = 0
    total_files = 0
    for path in dataset_root.rglob("*"):
        if not path.is_file():
            continue
        total_files += 1
        size = path.stat().st_size
        total_bytes += size
        extension_counts[path.suffix.lower() or "[no extension]"] += 1

    if preferred_sample and (
        not candidate
        or candidate["partition"] != preferred_partition
        or candidate["stem"] != preferred_stem
    ):
        raise FileNotFoundError(f"Representative sample '{preferred_sample}' was not found in the dataset labels.")

    total_samples = overall["label_files"]
    expected_keypoints = overall["instances"] * len(KEYPOINTS)
    summary = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "dataset_root": dataset_root.name,
        "layout": layout,
        "partitions": partition_rows,
        "overall": {
            "partitions": len(partition_rows),
            "samples": total_samples,
            "instances": overall["instances"],
            "keypoints": overall["keypoints"],
            "expected_keypoints": expected_keypoints,
            "mean_instances_per_frame": round(overall["instances"] / total_samples, 2) if total_samples else 0,
            "schema_complete_instances": overall["complete_groups"],
            "schema_completeness_percent": percent(overall["complete_groups"], overall["instances"]),
            "valid_boxes": overall["instances"] - overall["invalid_boxes"],
            "invalid_boxes": overall["invalid_boxes"],
            "unmatched_keypoints": overall["unmatched_points"],
            "fully_paired_samples": overall["fully_paired_samples"],
            "pairing_percent": percent(overall["fully_paired_samples"], total_samples),
            "rgb_files": overall["rgb_files"],
            "label_files": overall["label_files"],
            "depth_aligned_files": overall["depth_aligned_files"],
            "depth_aligned_pp_files": overall["depth_aligned_pp_files"],
            "depth_raw_files": overall["depth_raw_files"],
            "camera_units": len(camera_serials),
            "device_names": sorted(device_names),
            "calibration_mapping_records": len(calibration_mapping),
            "calibration_mapping_percent": percent(overall["calibration_mapped_samples"], total_samples)
            if calibration_mapping
            else None,
            "total_files": total_files,
            "total_bytes": total_bytes,
            "total_gb_decimal": round(total_bytes / 1_000_000_000, 2),
            "total_gib": round(total_bytes / (1024**3), 2),
        },
        "distributions": {
            "ripeness": dict(sorted(ripeness.items(), key=value_sort_key)),
            "occlusion": dict(sorted(occlusion.items(), key=value_sort_key)),
            "keypoint_visibility": dict(sorted(visibility.items(), key=value_sort_key)),
            "image_size": dict(image_sizes),
            "file_extension": dict(extension_counts),
        },
        "annotation_schema": {
            "instance_label": "strawberry",
            "keypoints": list(KEYPOINTS),
            "ripeness_values_observed": sorted(ripeness),
            "occlusion_values_observed": sorted(occlusion),
            "visibility_values_observed": sorted(visibility),
            "note": "Ripeness and occlusion are reported as stored numeric categories; no undocumented semantic names are inferred.",
        },
        "representative_sample": {
            "partition": candidate["partition"] if candidate else None,
            "stem": candidate["stem"] if candidate else None,
            "instances": candidate["instances"] if candidate else 0,
        },
    }
    return summary, candidate or {}


def fit_text(draw: ImageDraw.ImageDraw, text: str, max_width: int, start_size: int, bold: bool = False):
    size = start_size
    while size > 12:
        selected = font(size, bold=bold)
        if draw.textbbox((0, 0), text, font=selected)[2] <= max_width:
            return selected
        size -= 1
    return font(12, bold=bold)


def badge(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, fill: str, text_fill: str = "#ffffff"):
    selected = font(22, bold=True, mono=True)
    bbox = draw.textbbox((0, 0), text, font=selected)
    width = bbox[2] - bbox[0] + 24
    height = 36
    x, y = xy
    draw.rounded_rectangle((x, y, x + width, y + height), radius=4, fill=fill)
    draw.text((x + 12, y + 5), text, font=selected, fill=text_fill)
    return width


def render_annotation(dataset_root: Path, candidate: dict, output_path: Path) -> None:
    stem = candidate["stem"]
    rgb = Image.open(candidate["rgb_path"]).convert("RGB")
    record = candidate["record"]

    canvas = Image.new("RGB", (1800, 1080), COLORS["paper"])
    draw = ImageDraw.Draw(canvas)
    draw.rectangle((0, 0, 1800, 90), fill=COLORS["navy"])
    draw.text((56, 23), "INSTANCE ANNOTATION VIEW", font=font(21, bold=True, mono=True), fill="#c9d9e2")
    draw.text(
        (56, 116),
        f"{candidate['partition']} / {stem}",
        font=font(44, bold=True),
        fill=COLORS["navy"],
    )
    draw.text(
        (56, 172),
        "Bounding boxes, five anatomical landmarks, and stored categorical attributes",
        font=font(23),
        fill=COLORS["muted"],
    )

    image_x, image_y, image_w, image_h = 56, 230, 1260, 709
    rendered = rgb.resize((image_w, image_h), Image.Resampling.LANCZOS)
    canvas.paste(rendered, (image_x, image_y))
    draw.rectangle((image_x, image_y, image_x + image_w, image_y + image_h), outline=COLORS["navy"], width=2)

    sx = image_w / rgb.width
    sy = image_h / rgb.height
    boxes: dict[object, dict] = {}
    point_shapes: list[dict] = []
    for shape in record.get("shapes", []):
        if shape.get("shape_type") == "rectangle" or shape.get("label") == "strawberry":
            boxes[shape.get("group_id")] = shape
        elif shape.get("shape_type") == "point" or shape.get("label") in KEYPOINTS:
            point_shapes.append(shape)

    for group_id, shape in boxes.items():
        points = shape.get("points", [])
        if len(points) != 2:
            continue
        (x1, y1), (x2, y2) = points
        left = image_x + int(min(x1, x2) * sx)
        top = image_y + int(min(y1, y2) * sy)
        right = image_x + int(max(x1, x2) * sx)
        bottom = image_y + int(max(y1, y2) * sy)
        draw.rectangle(
            (left, top, right, bottom),
            outline=COLORS["red"],
            width=ANNOTATION_BOX_WIDTH,
        )
        code = f"#{group_id}  R{shape.get('ripeness')} · O{shape.get('occlusion')}"
        label_font = font(17, bold=True, mono=True)
        text_box = draw.textbbox((0, 0), code, font=label_font)
        label_w = text_box[2] - text_box[0] + 16
        label_y = max(image_y, top - 27)
        draw.rectangle((left, label_y, left + label_w, label_y + 27), fill=COLORS["red"])
        draw.text((left + 8, label_y + 3), code, font=label_font, fill="#ffffff")

    for shape in point_shapes:
        coords = shape.get("points", [])
        visibility = shape.get("visibility")
        if not coords or visibility == 0:
            continue
        x, y = coords[0]
        if x < 0 or y < 0:
            continue
        px = image_x + int(x * sx)
        py = image_y + int(y * sy)
        color = KEYPOINT_COLORS.get(shape.get("label"), COLORS["blue"])
        radius = ANNOTATION_KEYPOINT_RADIUS
        draw.ellipse((px - radius - 2, py - radius - 2, px + radius + 2, py + radius + 2), fill="#ffffff")
        if visibility == 1:
            draw.ellipse((px - radius, py - radius, px + radius, py + radius), outline=color, width=2)
        else:
            draw.ellipse((px - radius, py - radius, px + radius, py + radius), fill=color)

    panel_x = 1370
    draw.text((panel_x, 240), "Annotation key", font=font(29, bold=True), fill=COLORS["navy"])
    draw.line((panel_x, 285, 1744, 285), fill=COLORS["line"], width=2)
    draw.rectangle(
        (panel_x, 320, panel_x + 46, 350),
        outline=COLORS["red"],
        width=ANNOTATION_BOX_WIDTH,
    )
    draw.text((panel_x + 66, 317), "strawberry bbox", font=font(21), fill=COLORS["ink"])
    draw.text((panel_x, 382), "Landmarks", font=font(17, bold=True, mono=True), fill=COLORS["muted"])
    y = 425
    for name in KEYPOINTS:
        color = KEYPOINT_COLORS[name]
        draw.ellipse((panel_x, y + 4, panel_x + 18, y + 22), fill=color)
        draw.text((panel_x + 34, y), name, font=font(21), fill=COLORS["ink"])
        y += 48

    draw.text((panel_x, 690), "Attributes", font=font(17, bold=True, mono=True), fill=COLORS["muted"])
    badge(draw, (panel_x, 732), "R0 / R1", COLORS["navy"])
    draw.text((panel_x, 779), "ripeness category", font=font(18), fill=COLORS["muted"])
    badge(draw, (panel_x, 822), "O0 / O1", COLORS["blue"])
    draw.text((panel_x, 869), "occlusion category", font=font(18), fill=COLORS["muted"])

    draw.text((56, 972), f"{len(boxes)} instances in this representative frame", font=font(22, bold=True), fill=COLORS["navy"])
    draw.text(
        (56, 1010),
        "Filled keypoint: visibility=2   ·   outlined keypoint: visibility=1   ·   omitted: visibility=0",
        font=font(19),
        fill=COLORS["muted"],
    )
    canvas.save(output_path, "WEBP", quality=91, method=6)


def resolve_legacy_calibration(collection: Path, stem: str) -> tuple[dict, str]:
    mapping_path = collection / "camera_calibration_mapping.json"
    if mapping_path.exists():
        mapping = read_json(mapping_path)
        filename = mapping.get("prefix_to_calibration", {}).get(stem[:2])
        if filename:
            calibration = read_json(collection / filename)
            return calibration, filename
    default_path = collection / "camera_calibration.json"
    if default_path.exists():
        return read_json(default_path), default_path.name
    candidates = sorted(collection.glob("camera_calibration_*.json"))
    if candidates:
        return read_json(candidates[0]), candidates[0].name
    raise FileNotFoundError(f"No calibration found for {collection.name}/{stem}")


def render_pointcloud(
    dataset_root: Path,
    candidate: dict,
    output_path: Path,
    interactive_output_path: Path | None = None,
    max_foreground_depth: float = 0.4,
) -> dict:
    if not 0.15 < max_foreground_depth <= 2.5:
        raise ValueError("Maximum foreground depth must be greater than 0.15 m and no more than 2.5 m.")
    stem = candidate["stem"]
    rgb = np.asarray(Image.open(candidate["rgb_path"]).convert("RGB"))
    depth = np.asarray(Image.open(candidate["depth_pp_path"]))
    if candidate.get("calibration_path"):
        calibration_path = Path(candidate["calibration_path"])
        calibration = read_json(calibration_path)
        calibration_name = calibration_path.name
    else:
        calibration, calibration_name = resolve_legacy_calibration(candidate["legacy_calibration_root"], stem)
    intrinsics = calibration["intrinsics"]["color"]
    scale = float(calibration["depth_scale_m"])

    step = 2
    sampled_depth = depth[::step, ::step].astype(np.float32) * scale
    sampled_rgb = rgb[::step, ::step]
    vv, uu = np.mgrid[0 : depth.shape[0] : step, 0 : depth.shape[1] : step]
    base_valid = np.isfinite(sampled_depth) & (sampled_depth > 0.15) & (sampled_depth < 2.5)

    boxes = [
        shape["points"]
        for shape in candidate["record"].get("shapes", [])
        if (shape.get("shape_type") == "rectangle" or shape.get("label") == "strawberry")
        and len(shape.get("points", [])) == 2
    ]
    depth_min, depth_max = 0.15, max_foreground_depth
    roi_left, roi_top, roi_right, roi_bottom = 0, 0, depth.shape[1] - 1, depth.shape[0] - 1
    if boxes:
        x_values = [float(point[0]) for box in boxes for point in box]
        y_values = [float(point[1]) for box in boxes for point in box]
        subject_left, subject_right = min(x_values), max(x_values)
        subject_top, subject_bottom = min(y_values), max(y_values)
        subject_width = max(subject_right - subject_left, 1.0)
        subject_height = max(subject_bottom - subject_top, 1.0)
        roi_left = max(0, int(math.floor(subject_left - max(48.0, subject_width * 0.12))))
        roi_right = min(
            depth.shape[1] - 1,
            int(math.ceil(subject_right + max(48.0, subject_width * 0.12))),
        )
        roi_top = max(0, int(math.floor(subject_top - max(64.0, subject_height * 0.25))))
        roi_bottom = min(
            depth.shape[0] - 1,
            int(math.ceil(subject_bottom + max(56.0, subject_height * 0.18))),
        )

        subject_depth_values: list[np.ndarray] = []
        for box in boxes:
            (x1, y1), (x2, y2) = box
            left = max(0, int(math.floor(min(x1, x2))))
            right = min(depth.shape[1], int(math.ceil(max(x1, x2))) + 1)
            top = max(0, int(math.floor(min(y1, y2))))
            bottom = min(depth.shape[0], int(math.ceil(max(y1, y2))) + 1)
            box_depth = depth[top:bottom, left:right].astype(np.float32) * scale
            subject_depth_values.append(
                box_depth[np.isfinite(box_depth) & (box_depth > 0.15) & (box_depth < 2.5)]
            )
        usable_subject_depth = [values for values in subject_depth_values if values.size]
        if usable_subject_depth:
            subject_depth = np.concatenate(usable_subject_depth)
            subject_near, subject_far = np.percentile(subject_depth, [2, 98])
            depth_min = max(0.15, float(subject_near) - 0.08)
            depth_max = min(max_foreground_depth, float(subject_far) + 0.22)

    foreground_roi = (
        (uu >= roi_left)
        & (uu <= roi_right)
        & (vv >= roi_top)
        & (vv <= roi_bottom)
        & (sampled_depth >= depth_min)
        & (sampled_depth <= depth_max)
    )
    valid = base_valid & foreground_roi
    if not np.any(valid):
        raise RuntimeError(f"Foreground filtering removed every point for {candidate['partition']}/{stem}.")
    z = sampled_depth[valid]
    u = uu[valid].astype(np.float32)
    v = vv[valid].astype(np.float32)
    color = sampled_rgb[valid]
    x = (u - float(intrinsics["ppx"])) / float(intrinsics["fx"]) * z
    y = -(v - float(intrinsics["ppy"])) / float(intrinsics["fy"]) * z

    points = np.column_stack((x, y, z))
    center = np.median(points, axis=0)
    centered = points - center
    if interactive_output_path is not None:
        target_count = min(120000, len(centered))
        sample_indexes = np.linspace(0, len(centered) - 1, target_count, dtype=np.int32)
        interactive_points = centered[sample_indexes]
        extent = max(float(np.max(np.ptp(interactive_points, axis=0))), 1e-6)
        interactive_points = interactive_points / extent
        interactive_colors = color[sample_indexes]
        interactive_payload = {
            "version": 1,
            "sample": f"{candidate['partition']}/{stem}",
            "point_count": int(target_count),
            "positions": np.round(interactive_points.astype(np.float64), 5).reshape(-1).tolist(),
            "colors": interactive_colors.astype(np.uint8).reshape(-1).tolist(),
        }
        interactive_output_path.write_text(
            json.dumps(interactive_payload, separators=(",", ":")),
            encoding="utf-8",
        )
    yaw = math.radians(0)
    pitch = math.radians(0)
    rotation_y = np.array(
        [[math.cos(yaw), 0, math.sin(yaw)], [0, 1, 0], [-math.sin(yaw), 0, math.cos(yaw)]],
        dtype=np.float32,
    )
    rotation_x = np.array(
        [[1, 0, 0], [0, math.cos(pitch), -math.sin(pitch)], [0, math.sin(pitch), math.cos(pitch)]],
        dtype=np.float32,
    )
    rotated = centered @ rotation_y.T @ rotation_x.T

    width, height = 1800, 1120
    plot_left, plot_top, plot_right, plot_bottom = 54, 174, 1330, 1044
    plot_w = plot_right - plot_left
    plot_h = plot_bottom - plot_top
    camera_distance = max(float(np.ptp(rotated[:, 2])) * 3.0, 1.2)
    denom = rotated[:, 2] + camera_distance
    raw_projected_x = rotated[:, 0] / denom
    raw_projected_y = -rotated[:, 1] / denom
    qx1, qx2 = np.percentile(raw_projected_x, [0.5, 99.5])
    qy1, qy2 = np.percentile(raw_projected_y, [0.5, 99.5])
    fit_scale = min(
        plot_w * 0.88 / max(float(qx2 - qx1), 1e-6),
        plot_h * 0.82 / max(float(qy2 - qy1), 1e-6),
    )
    projected_x = plot_left + plot_w / 2 + fit_scale * (raw_projected_x - (qx1 + qx2) / 2)
    projected_y = plot_top + plot_h / 2 + fit_scale * (raw_projected_y - (qy1 + qy2) / 2)
    inside = (
        (projected_x >= plot_left)
        & (projected_x < plot_right)
        & (projected_y >= plot_top)
        & (projected_y < plot_bottom)
        & (denom > 0)
    )
    projected_x = projected_x[inside].astype(np.int32)
    projected_y = projected_y[inside].astype(np.int32)
    denom = denom[inside]
    color = color[inside]

    canvas = Image.new("RGB", (width, height), "#081521")
    draw = ImageDraw.Draw(canvas)
    draw.rectangle((0, 0, width, 104), fill="#0d2332")
    draw.text((54, 26), "RGB-COLORED 3D POINT CLOUD", font=font(24, bold=True, mono=True), fill="#bcd0db")
    draw.text((54, 118), f"{candidate['partition']} / {stem}", font=font(32, bold=True), fill="#f2f6f8")

    for gx in range(plot_left, plot_right + 1, 128):
        draw.line((gx, plot_top, gx, plot_bottom), fill="#163142", width=1)
    for gy in range(plot_top, plot_bottom + 1, 112):
        draw.line((plot_left, gy, plot_right, gy), fill="#163142", width=1)
    draw.rectangle((plot_left, plot_top, plot_right, plot_bottom), outline="#335064", width=2)

    order = np.argsort(denom)[::-1]
    for idx in order:
        px = int(projected_x[idx])
        py = int(projected_y[idx])
        r, g, b = (int(channel) for channel in color[idx])
        boost = 0.84 + 0.16 * (1.0 - min(max((denom[idx] - denom.min()) / (np.ptp(denom) + 1e-6), 0), 1))
        fill = (int(r * boost), int(g * boost), int(b * boost))
        draw.rectangle((px, py, px + 2, py + 2), fill=fill)

    panel_left = 1382
    draw.text((panel_left, 174), "Filtered depth", font=font(27, bold=True), fill="#f2f6f8")
    draw.line((panel_left, 218, 1745, 218), fill="#335064", width=2)
    valid_full = depth > 0
    p10, p50, p90 = np.percentile(z, [10, 50, 90]) if z.size else (0, 0, 0)
    kept_sampled_percent = percent(int(valid.sum()), int(base_valid.sum()))
    metrics = [
        ("Kept samples", f"{kept_sampled_percent:.1f}%"),
        ("P10 depth", f"{p10:.2f} m"),
        ("Median depth", f"{p50:.2f} m"),
        ("P90 depth", f"{p90:.2f} m"),
        ("Rendered points", f"{len(order):,}"),
    ]
    y = 258
    for label, value in metrics:
        draw.text((panel_left, y), label.upper(), font=font(15, bold=True, mono=True), fill="#82a0b2")
        draw.text((panel_left, y + 27), value, font=font(30, bold=True), fill="#f2f6f8")
        draw.line((panel_left, y + 72, 1745, y + 72), fill="#243e50", width=1)
        y += 106

    draw.text((panel_left, 812), "Projection", font=font(17, bold=True, mono=True), fill="#82a0b2")
    draw.text((panel_left, 847), "Label-aware foreground crop", font=font(19), fill="#dbe7ed")
    draw.text((panel_left, 882), "Color-camera intrinsics", font=font(19), fill="#dbe7ed")
    draw.text((panel_left, 917), calibration_name, font=font(16, mono=True), fill="#82a0b2")

    origin = (1450, 1021)
    draw.line((origin[0], origin[1], origin[0] + 95, origin[1]), fill="#e06b72", width=4)
    draw.line((origin[0], origin[1], origin[0], origin[1] - 76), fill="#6fb886", width=4)
    draw.line((origin[0], origin[1], origin[0] + 55, origin[1] - 48), fill="#6f9fc0", width=4)
    draw.text((origin[0] + 104, origin[1] - 12), "X", font=font(16, bold=True, mono=True), fill="#e06b72")
    draw.text((origin[0] - 7, origin[1] - 100), "Y", font=font(16, bold=True, mono=True), fill="#6fb886")
    draw.text((origin[0] + 61, origin[1] - 68), "Z", font=font(16, bold=True, mono=True), fill="#6f9fc0")

    canvas.save(output_path, "WEBP", quality=91, method=6)
    return {
        "valid_pixel_percent": percent(int(valid_full.sum()), depth.size),
        "p10_depth_m": round(float(p10), 3),
        "median_depth_m": round(float(p50), 3),
        "p90_depth_m": round(float(p90), 3),
        "rendered_points": int(len(order)),
        "foreground_depth_min_m": round(depth_min, 3),
        "foreground_depth_max_m": round(depth_max, 3),
        "foreground_roi": [roi_left, roi_top, roi_right, roi_bottom],
        "kept_sampled_pixels_percent": kept_sampled_percent,
        "calibration_file": calibration_name,
    }


def panel(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], label: str, title: str):
    draw.rounded_rectangle(box, radius=6, fill=COLORS["paper"], outline=COLORS["line"], width=2)
    x1, y1, _, _ = box
    draw.text((x1 + 30, y1 + 25), label, font=font(16, bold=True, mono=True), fill=COLORS["blue"])
    draw.text((x1 + 30, y1 + 58), title, font=font(27, bold=True), fill=COLORS["navy"])


def draw_horizontal_bars(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    labels: list[str],
    values: list[int],
    color: str,
):
    x1, y1, x2, y2 = box
    graph_left = x1 + 156
    graph_right = x2 - 80
    graph_top = y1 + 118
    graph_bottom = y2 - 34
    max_value = max(values) if values else 1
    row_h = (graph_bottom - graph_top) / max(len(values), 1)
    for index, (label, value) in enumerate(zip(labels, values)):
        y = graph_top + index * row_h
        bar_h = max(10, int(row_h * 0.52))
        draw.text((x1 + 30, y - 2), label, font=font(16, mono=True), fill=COLORS["muted"])
        width = int((graph_right - graph_left) * value / max_value)
        draw.rounded_rectangle(
            (graph_left, y, graph_left + width, y + bar_h),
            radius=min(5, bar_h // 2),
            fill=color,
        )
        draw.text((graph_left + width + 10, y - 3), f"{value:,}", font=font(16, bold=True), fill=COLORS["ink"])


def draw_category_bars(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    distribution: dict[str, int],
    color: str,
):
    x1, y1, x2, y2 = box
    items = list(distribution.items())
    total = sum(distribution.values())
    graph_left = x1 + 48
    graph_right = x2 - 48
    graph_top = y1 + 132
    graph_bottom = y2 - 66
    max_value = max((value for _, value in items), default=1)
    gap = 42
    bar_w = int((graph_right - graph_left - gap * (len(items) - 1)) / max(len(items), 1))
    for index, (label, value) in enumerate(items):
        left = graph_left + index * (bar_w + gap)
        height = int((graph_bottom - graph_top) * value / max_value)
        top = graph_bottom - height
        draw.rounded_rectangle((left, top, left + bar_w, graph_bottom), radius=5, fill=color)
        draw.text((left, graph_bottom + 12), label, font=font(18, bold=True, mono=True), fill=COLORS["navy"])
        label_y = max(y1 + 105, top - 28)
        draw.text(
            (left, label_y),
            f"{value:,} · {percent(value, total):.1f}%",
            font=font(16, bold=True),
            fill=COLORS["ink"],
        )


def render_distributions(summary: dict, output_path: Path) -> None:
    canvas = Image.new("RGB", (1800, 1250), COLORS["panel"])
    draw = ImageDraw.Draw(canvas)
    draw.rectangle((0, 0, 1800, 112), fill=COLORS["navy"])
    draw.text((52, 25), "DATASET COMPOSITION", font=font(25, bold=True, mono=True), fill="#c9d9e2")
    overall = summary["overall"]
    subtitle = (
        f"{overall['samples']:,} synchronized RGB-D frames · "
        f"{overall['instances']:,} strawberry instances · "
        f"{overall['keypoints']:,} keypoints"
    )
    draw.text((52, 132), subtitle, font=font(28, bold=True), fill=COLORS["navy"])
    draw.text(
        (52, 177),
        "Split: 70% training / 15% validation / 15% test."
        if summary.get("split_update")
        else "Counts are computed directly from the current data/ directory.",
        font=font(20),
        fill=COLORS["muted"],
    )

    boxes = {
        "samples": (44, 232, 884, 744),
        "instances": (916, 232, 1756, 744),
        "ripeness": (44, 776, 608, 1198),
        "occlusion": (628, 776, 1192, 1198),
        "visibility": (1212, 776, 1756, 1198),
    }
    grouping = "split" if summary["layout"] in ("train_val", "train_val_test") else "collection"
    panel(draw, boxes["samples"], "A", f"Frames by {grouping}")
    panel(draw, boxes["instances"], "B", f"Instances by {grouping}")
    panel(draw, boxes["ripeness"], "C", "Ripeness categories")
    panel(draw, boxes["occlusion"], "D", "Occlusion categories")
    panel(draw, boxes["visibility"], "E", "Keypoint visibility")

    rows = summary["partitions"]
    labels = [row["label"] for row in rows]
    draw_horizontal_bars(draw, boxes["samples"], labels, [row["samples"] for row in rows], COLORS["blue"])
    if all(row.get("instances") is not None for row in rows):
        draw_horizontal_bars(draw, boxes["instances"], labels, [row["instances"] for row in rows], COLORS["red"])
    else:
        x1, y1, _, _ = boxes["instances"]
        for index, label in enumerate(labels):
            draw.text((x1 + 30, y1 + 118 + index * 120), label, font=font(16, mono=True), fill=COLORS["muted"])
    draw_category_bars(draw, boxes["ripeness"], summary["distributions"]["ripeness"], COLORS["green"])
    draw_category_bars(draw, boxes["occlusion"], summary["distributions"]["occlusion"], COLORS["gold"])
    draw_category_bars(draw, boxes["visibility"], summary["distributions"]["keypoint_visibility"], COLORS["blue"])
    canvas.save(output_path, "WEBP", quality=92, method=6)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("dataset_root", type=Path)
    parser.add_argument("output_dir", type=Path)
    parser.add_argument(
        "--representative-sample",
        default="train/0238",
        help="Sample used for the annotation and point-cloud figures (default: train/0238).",
    )
    parser.add_argument(
        "--foreground-depth-max",
        type=float,
        default=0.4,
        help="Maximum depth retained in the foreground point cloud, in meters (default: 0.4).",
    )
    args = parser.parse_args()
    dataset_root = args.dataset_root.resolve()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    summary, candidate = scan_dataset(dataset_root, args.representative_sample)
    if not candidate:
        raise RuntimeError("No labeled representative sample could be selected.")

    annotation_path = output_dir / "dataset-annotation-overview.webp"
    distribution_path = output_dir / "dataset-distributions.webp"
    pointcloud_path = output_dir / "dataset-pointcloud.webp"
    interactive_pointcloud_path = output_dir / "dataset-pointcloud-data.json"
    render_annotation(dataset_root, candidate, annotation_path)
    summary["representative_sample"]["depth"] = render_pointcloud(
        dataset_root,
        candidate,
        pointcloud_path,
        interactive_pointcloud_path,
        args.foreground_depth_max,
    )
    render_distributions(summary, distribution_path)

    summary_path = output_dir / "dataset-summary.json"
    summary_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(summary["overall"], indent=2, ensure_ascii=False))
    print(f"Representative sample: {candidate['partition']}/{candidate['stem']}")
    print(f"Wrote: {summary_path}")
    print(f"Wrote: {annotation_path}")
    print(f"Wrote: {distribution_path}")
    print(f"Wrote: {pointcloud_path}")
    print(f"Wrote: {interactive_pointcloud_path}")


if __name__ == "__main__":
    main()
